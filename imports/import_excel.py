import sqlite3
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = BASE_DIR / "historische_verbraeuche.xlsx"
DATABASE_PATH = BASE_DIR / "energy_cost_tracker.db"
SHEET_INDEX = 0
SOURCE = "manual_import"  # Quelle für die importierten Daten


def import_excel():
    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True,
    )
    sheet = workbook.worksheets[SHEET_INDEX]

    rows = sheet.iter_rows(min_row=2, values_only=True)

    imported = 0
    skipped = 0
    errors = []

    try:
        with sqlite3.connect(DATABASE_PATH) as connection:
            for row_number, row in enumerate(rows, start=2):
                start, end, consumption, feed_in, delta, note = row

                # Leere Zeilen überspringen
                if all(value is None for value in row):
                    skipped += 1
                    continue

                # Zeitstempel müssen vorhanden sein
                if not start or not end:
                    errors.append(f"Zeile {row_number}: Zeitstempel fehlen")
                    continue

                # Zahlenwerte müssen gültig sein
                try:
                    consumption = float(consumption)
                    feed_in = float(feed_in)
                    delta = float(delta)
                except TypeError, ValueError:
                    errors.append(f"Zeile {row_number}: Ungültiger Zahlenwert")
                    continue

                cursor = connection.execute(
                    """
                    INSERT OR IGNORE INTO meter_intervals (
                        interval_start_utc,
                        interval_end_utc,
                        consumption_meter_kwh,
                        feed_in_meter_kwh,
                        consumption_delta_kwh,
                        note,
                        source
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        str(start),
                        str(end),
                        consumption,
                        feed_in,
                        delta,
                        str(note) if note is not None else None,
                        SOURCE,
                    ),
                )

                if cursor.rowcount == 1:
                    imported += 1
                else:
                    skipped += 1

    finally:
        workbook.close()

    print(f"Neu importiert: {imported}")
    print(f"Übersprungen (leer oder bereits vorhanden): {skipped}")
    print(f"Fehlerhafte Zeilen: {len(errors)}")

    for error in errors[:20]:
        print(error)

    if len(errors) > 20:
        print(f"... und {len(errors) - 20} weitere Fehler.")


if __name__ == "__main__":
    import_excel()
