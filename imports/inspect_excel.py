"""Validate historical meter readings from an Excel workbook."""

from dataclasses import dataclass
from datetime import UTC, datetime
import math
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

EXPECTED_HEADERS = (
    "Von (UTC nach ISO 8601)",
    "Bis (UTC nach ISO 8601)",
    "Zählerstand - Verbrauch (kWh)",
    "Zählerstand - Einspeisung (kWh)",
    "Delta (Verbrauch) in kWh",
    "Anmerkung",
)


@dataclass(frozen=True, slots=True)
class ImportedInterval:
    """One validated interval from an uploaded workbook."""

    start: str
    end: str
    consumption: float
    feed_in: float | None
    delta: float
    note: str | None


@dataclass(frozen=True, slots=True)
class ExcelInspection:
    """Validated workbook contents and row-level errors."""

    rows: list[ImportedInterval]
    empty_rows: int
    errors: list[str]
    error_count: int


class InvalidExcelFile(Exception):
    """The uploaded file cannot be read as an Excel workbook."""


def _parse_timestamp(value: object) -> datetime:
    """Parse an Excel or ISO 8601 timestamp and normalize it to UTC."""
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        parsed = datetime.fromisoformat(value.strip())
    else:
        raise TypeError("Zeitstempel fehlt oder ist ungültig")

    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed.astimezone(UTC)


def normalize_timestamp(value: str) -> str:
    """Normalize a stored timestamp for comparisons across ISO representations."""
    return _parse_timestamp(value).isoformat()


def _parse_number(value: object, column: str) -> float:
    """Parse a finite numeric cell and include its source value in errors."""
    if value is None or isinstance(value, bool):
        raise ValueError(f"{column}: ungültiger Zahlenwert {value!r}")
    try:
        result = float(value)
    except (TypeError, ValueError) as err:
        raise ValueError(f"{column}: ungültiger Zahlenwert {value!r}") from err
    if not math.isfinite(result):
        raise ValueError(f"{column}: ungültiger Zahlenwert {value!r}")
    return result


def _parse_optional_number(value: object, column: str) -> float | None:
    """Parse an optional finite numeric cell."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return _parse_number(value, column)


def _parse_row(row: tuple[object, ...]) -> ImportedInterval:
    """Validate and normalize one data row."""
    if any(value is not None for value in row[len(EXPECTED_HEADERS) :]):
        raise ValueError("Unerwartete zusätzliche Spalten")
    values = tuple(row[: len(EXPECTED_HEADERS)]) + (None,) * max(
        0, len(EXPECTED_HEADERS) - len(row)
    )
    start = _parse_timestamp(values[0])
    end = _parse_timestamp(values[1])
    if end <= start:
        raise ValueError("Intervallende muss nach dem Intervallbeginn liegen")
    return ImportedInterval(
        start=start.isoformat(),
        end=end.isoformat(),
        consumption=_parse_number(values[2], EXPECTED_HEADERS[2]),
        feed_in=_parse_optional_number(values[3], EXPECTED_HEADERS[3]),
        delta=_parse_number(values[4], EXPECTED_HEADERS[4]),
        note=str(values[5]) if values[5] is not None else None,
    )


def _get_sheet(workbook: Any) -> Any:
    """Get the first sheet after verifying its column headers."""
    if not workbook.worksheets:
        raise InvalidExcelFile
    sheet = workbook.worksheets[0]
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    normalized_header_values = [
        str(value).strip() if value is not None else "" for value in header
    ]
    while normalized_header_values and not normalized_header_values[-1]:
        normalized_header_values.pop()
    if tuple(normalized_header_values) != EXPECTED_HEADERS:
        raise InvalidExcelFile
    return sheet


def inspect_excel(file_path: Path) -> ExcelInspection:
    """Validate the workbook without changing the database."""
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError) as err:
        raise InvalidExcelFile from err

    try:
        try:
            sheet = _get_sheet(workbook)
        except (
            BadZipFile,
            InvalidFileException,
            OSError,
            ParseError,
            ValueError,
        ) as err:
            raise InvalidExcelFile from err
        rows: list[ImportedInterval] = []
        errors: list[str] = []
        error_count = 0
        empty_rows = 0
        try:
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=2, values_only=True), start=2
            ):
                if all(value is None for value in row):
                    empty_rows += 1
                    continue
                try:
                    rows.append(_parse_row(row))
                except (OverflowError, TypeError, ValueError) as err:
                    error_count += 1
                    if len(errors) < 20:
                        errors.append(f"Zeile {row_number}: {err}")
        except (
            BadZipFile,
            InvalidFileException,
            OSError,
            ParseError,
            ValueError,
        ) as err:
            raise InvalidExcelFile from err
    finally:
        workbook.close()

    if not rows and not errors:
        errors.append("Die Tabelle enthält keine Messwerte.")
        error_count = 1
    return ExcelInspection(
        rows=rows, empty_rows=empty_rows, errors=errors, error_count=error_count
    )
